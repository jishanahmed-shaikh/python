import mysql.connector
from docx import Document
from docx2pdf import convert
import os

def gen_report(req_id):
    # Connect to your MySQL database
    mydb = mysql.connector.connect(
        host="localhost",
        user="root",
        password="root",
        database="bloodbridge"
    )

    # Retrieve the data from your MySQL table
    mycursor = mydb.cursor()
    mycursor.execute("SELECT d.username, d.units, b.name, d.appointment, d.timing FROM donation_requests d,blood_banks b where b.Bid=d.Bid and d.request_id=%s",(req_id,))
    data = mycursor.fetchone()
    name=data[0]

    # Open your Word document and locate the placeholders
    document = Document("certificate_template.docx")
    placeholders = {
        "<username>": data[0],
        "<units>": str(data[1]),
        "<blood bank name>": data[2],
        "<appointment>": str(data[3]),
        "<timing>": str(data[4])
    }

    # Replace the placeholders with the values retrieved from MySQL
    for p in document.paragraphs:
        for key in placeholders:
            if key in p.text:
                p.text = p.text.replace(key, placeholders[key])

    # Save your updated Word document
    document.save("donation_report.docx")


    # Convert the Word document to PDF
    convert("donation_report.docx")

    # Move the PDF file to the reports directory
    if not os.path.exists("reports"):
        os.mkdir("reports")
    os.rename("donation_report.pdf", "reports/donation_report.pdf")

    # Delete the Word document
    os.remove("donation_report.docx")